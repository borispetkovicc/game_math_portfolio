"""Generate docs/Kastel_Fortune_math_model.xlsx — the Kastel Fortune math model as a LIVE Excel workbook.

Everything that can be exact in closed form is a formula: reel-strip counts (COUNTIF), per-reel symbol
probabilities, the single-line expected value, the scatter-count distribution (Poisson-binomial DP), the
free-spins Markov chain (mean recursion over the state grid), the RTP decomposition, a what-if scenario and
a bet-level table. Quantities that need the full enumeration of the grid (hit frequency, variance, tail
probabilities) are pasted from results/kastel_exact.json and labelled as such.

Why the closed form is exact for this game: wilds sit on reels 2-4 only, so the paying symbol of a line is
always the reel-1 symbol and P(run >= k) = p1(s) * prod_{i=2..k} (p_i(s) + p_i(W)).  The scatter formula
P(visible) = rows * count / length is exact because each strip carries exactly one scatter.

Usage: python scripts/make_excel_model.py && python <xlsx-skill>/scripts/recalc.py docs/Kastel_Fortune_math_model.xlsx
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from games import kastel_fortune as kf      # noqa: E402

OUT = ROOT / "docs" / "Kastel_Fortune_math_model.xlsx"

# ---- styles --------------------------------------------------------------
FONT = "Arial"
F_IN = Font(name=FONT, color="0000FF")                 # hard-coded input
F_LINK = Font(name=FONT, color="008000")               # link to another sheet
F_TXT = Font(name=FONT)
F_B = Font(name=FONT, bold=True)
F_T = Font(name=FONT, bold=True, size=14)
F_H = Font(name=FONT, bold=True, color="FFFFFF")
FILL_H = PatternFill("solid", fgColor="1F3864")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_REF = PatternFill("solid", fgColor="E7E6E6")
FILL_GRID = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PCT4 = "0.0000%"
PCT2 = "0.00%"
NUM4 = "0.0000"
NUM6 = "0.000000"
SCI = "0.000E+00"


def put(ws, ref, value, font=F_TXT, fmt=None, fill=None, bold=False, comment=None, align=None):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, bold=bold, color=font.color) if bold else font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if comment:
        c.comment = Comment(comment, "model")
    if align:
        c.alignment = Alignment(horizontal=align)
    return c


def header(ws, row, col0, labels):
    for j, t in enumerate(labels):
        c = ws.cell(row=row, column=col0 + j, value=t)
        c.font = F_H; c.fill = FILL_H; c.alignment = Alignment(horizontal="center", wrap_text=True); c.border = BOX


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# ---- model data ------------------------------------------------------------
cfg = kf.load_config()
SYM = cfg.symbols                      # W S H1 H2 H3 H4 A K Q J
PAYING = [s for i, s in enumerate(SYM) if cfg.paytable[i, 3:].sum() > 0]
N_REELS, ROWS, N_LINES = cfg.n_reels, cfg.rows, cfg.n_lines
exact = json.loads((ROOT / "results" / "kastel_exact.json").read_text())
rules = kf.RULES

wb = Workbook()
wb.calculation = CalcProperties(fullCalcOnLoad=True)

# =============================================================================
# Paytable
# =============================================================================
ws = wb.active; ws.title = "Paytable"
put(ws, "A1", "Kastel Fortune — paytable and game parameters", F_T)
put(ws, "A2", "Pays in multiples of the TOTAL bet. Blue = input. Line pays left to right from reel 1; wild substitutes for all but the scatter.")
header(ws, 4, 1, ["Symbol", "Name", "3 of a kind", "4 of a kind", "5 of a kind"])
PAY = {}                                                   # symbol -> {k: cell ref}
for i, s in enumerate(SYM):
    r = 5 + i
    put(ws, f"A{r}", s, F_B); put(ws, f"B{r}", kf.SYMBOL_NAMES[s])
    PAY[s] = {}
    for k in (3, 4, 5):
        col = L(k)                                         # C, D, E
        put(ws, f"{col}{r}", float(cfg.paytable[i, k]), F_IN, "0.0")
        PAY[s][k] = f"Paytable!${col}${r}"
    if s == "W":
        ws[f"B{r}"].comment = Comment("Wild has no pays of its own: it only substitutes. Wilds sit on reels 2-4 and a line run "
                                      "must start on reel 1, so a pure-wild run can never occur.", "model")
    if s == "S":
        ws[f"B{r}"].comment = Comment("Scatter pays anywhere on the grid (see below) and triggers free spins; it never pays on a line.", "model")
SYM_ROW = {s: 5 + i for i, s in enumerate(SYM)}
r0 = 5 + len(SYM) + 2
put(ws, f"A{r0}", "Scatter pays (anywhere on the grid)", F_B)
header(ws, r0 + 1, 1, ["Scatters", "Pay (x bet)", "Free spins awarded"])
SC_PAY, SC_SPINS = {}, {}
for j, k in enumerate((3, 4, 5)):
    r = r0 + 2 + j
    put(ws, f"A{r}", k, F_B); put(ws, f"B{r}", float(cfg.scatter_pays[k]), F_IN, "0.0"); put(ws, f"C{r}", rules.spins_by_trigger[k], F_IN)
    SC_PAY[k] = f"Paytable!$B${r}"; SC_SPINS[k] = f"Paytable!$C${r}"
r1 = r0 + 7
put(ws, f"A{r1}", "Game parameters", F_B)
params = [("Reels", N_REELS), ("Rows", ROWS), ("Paylines", N_LINES), ("Max win cap (x bet)", float(cfg.win_cap)),
          ("Retrigger: scatters needed", rules.retrigger_min), ("Retrigger: extra spins", rules.retrigger_spins),
          ("Max retriggers", rules.max_retriggers), ("Multiplier start", rules.mult_start),
          ("Multiplier step per winning spin", rules.mult_step), ("Multiplier cap", rules.mult_cap)]
P = {}
for j, (name, v) in enumerate(params):
    r = r1 + 1 + j
    put(ws, f"A{r}", name); put(ws, f"B{r}", v, F_IN)
    P[name] = f"Paytable!$B${r}"
ws[f"A{r1 + 6}"].comment = Comment("The free-spins grid on sheet FreeSpins is laid out for +5 spins per retrigger, max 4 retriggers, "
                                   "multiplier 1->5 in steps of 1. Changing these needs the grid regenerated (scripts/make_excel_model.py).", "model")
widths(ws, {"A": 30, "B": 22, "C": 14, "D": 14, "E": 18})

# =============================================================================
# Reels sheets
# =============================================================================
STRIP_ROWS = 50                                            # room for edits


def reels_sheet(name, reels, title):
    ws = wb.create_sheet(name)
    put(ws, "A1", title, F_T)
    put(ws, "A2", "Strips top-to-bottom (wrapping). Blue cells are inputs: edit a symbol and every probability, the line EV and the RTP recalculate.")
    header(ws, 4, 1, ["Pos"] + [f"Reel {i+1}" for i in range(N_REELS)])
    for p in range(STRIP_ROWS):
        put(ws, f"A{5+p}", p + 1)
    for i, strip in enumerate(reels):
        col = L(2 + i)
        for p, sid in enumerate(strip):
            put(ws, f"{col}{5+p}", SYM[int(sid)], F_IN, align="center")
    first, last = 5, 5 + STRIP_ROWS - 1
    # counts table
    c0 = 2 + N_REELS + 2                                   # column index of "Symbol"
    put(ws, f"{L(c0)}3", "Symbol counts per reel (COUNTIF)", F_B)
    header(ws, 4, c0, ["Symbol"] + [f"Reel {i+1}" for i in range(N_REELS)])
    CNT, PRB = {}, {}
    for j, s in enumerate(SYM):
        r = 5 + j
        put(ws, f"{L(c0)}{r}", s, F_B)
        CNT[s] = {}
        for i in range(N_REELS):
            sc = L(2 + i); cc = L(c0 + 1 + i)
            put(ws, f"{cc}{r}", f'=COUNTIF({sc}${first}:{sc}${last},${L(c0)}{r})')
            CNT[s][i] = f"{name}!${cc}${r}"
    rl = 5 + len(SYM)
    put(ws, f"{L(c0)}{rl}", "Strip length", F_B)
    LEN = {}
    for i in range(N_REELS):
        sc = L(2 + i); cc = L(c0 + 1 + i)
        put(ws, f"{cc}{rl}", f"=COUNTA({sc}${first}:{sc}${last})", bold=True)
        LEN[i] = f"{name}!${cc}${rl}"
    # probabilities table
    p0 = c0 + N_REELS + 2
    put(ws, f"{L(p0)}3", "P(symbol on a given row of the reel) = count / length", F_B)
    header(ws, 4, p0, ["Symbol"] + [f"Reel {i+1}" for i in range(N_REELS)])
    for j, s in enumerate(SYM):
        r = 5 + j
        put(ws, f"{L(p0)}{r}", s, F_B)
        PRB[s] = {}
        for i in range(N_REELS):
            cc = L(c0 + 1 + i); pc = L(p0 + 1 + i)
            put(ws, f"{pc}{r}", f"={cc}{r}/{cc}${rl}", fmt=PCT2)
            PRB[s][i] = f"{name}!${pc}${r}"
    put(ws, f"{L(p0)}{rl}", "check: sums to 1", F_B)
    for i in range(N_REELS):
        pc = L(p0 + 1 + i)
        put(ws, f"{pc}{rl}", f"=SUM({pc}5:{pc}{rl-1})", fmt="0.0000")
    widths(ws, {"A": 6, **{L(2 + i): 8 for i in range(N_REELS)}, L(c0): 14, L(p0): 14})
    return CNT, PRB, LEN


CNT_B, PRB_B, LEN_B = reels_sheet("Reels_Base", cfg.base_reels, "Kastel Fortune — base-game reel strips")
CNT_F, PRB_F, LEN_F = reels_sheet("Reels_FS", cfg.fs_reels, "Kastel Fortune — free-spins reel strips")

# =============================================================================
# Lines sheets (closed-form single-line EV)
# =============================================================================


def lines_block(ws, top, prb, title, note_ref_sheet):
    """Write the closed-form line block starting at row `top`. prb[s][i] -> cell ref of P(symbol s on reel i).
    Returns (cell of E[one line], cell of P(line pays), cell of validity check)."""
    put(ws, f"A{top}", title, F_B)
    put(ws, f"A{top+1}", "P(run >= k) = p1(s) x prod_(i=2..k) [p_i(s) + p_i(W)] — exact because reel 1 carries no wild, so the paying "
                         "symbol is always the reel-1 symbol. P(run = k) = P(run >= k) - P(run >= k+1).")
    hdr = ["Symbol"] + [f"p{i+1}(s)" for i in range(N_REELS)] + [f"p{i+1}(W)" for i in range(1, N_REELS)] \
        + [f"P(run>={k})" for k in range(1, 6)] + ["P(run=3)", "P(run=4)", "P(run=5)", "pay 3", "pay 4", "pay 5", "E[win] x bet", "P(pays)"]
    header(ws, top + 2, 1, hdr)
    r = top + 3
    cP = {i: L(2 + i) for i in range(N_REELS)}                       # p_i(s) columns B..F
    cW = {i: L(2 + N_REELS + (i - 1)) for i in range(1, N_REELS)}   # p_i(W) columns G..J
    cGE = {k: L(2 + N_REELS + (N_REELS - 1) + (k - 1)) for k in range(1, 6)}   # K..O
    cEQ = {k: L(2 + N_REELS + (N_REELS - 1) + 5 + (k - 3)) for k in (3, 4, 5)}  # P..R
    cPay = {k: L(2 + N_REELS + (N_REELS - 1) + 8 + (k - 3)) for k in (3, 4, 5)}  # S..U
    cE = L(2 + N_REELS + (N_REELS - 1) + 11)                        # V
    cH = L(2 + N_REELS + (N_REELS - 1) + 12)                        # W
    for s in PAYING:
        put(ws, f"A{r}", s, F_B)
        for i in range(N_REELS):
            put(ws, f"{cP[i]}{r}", f"={prb[s][i]}", F_LINK, PCT2)
        for i in range(1, N_REELS):
            put(ws, f"{cW[i]}{r}", f"={prb['W'][i]}", F_LINK, PCT2)
        put(ws, f"{cGE[1]}{r}", f"={cP[0]}{r}", fmt=NUM6)
        for k in range(2, 6):
            put(ws, f"{cGE[k]}{r}", f"={cGE[k-1]}{r}*({cP[k-1]}{r}+{cW[k-1]}{r})", fmt=NUM6)
        put(ws, f"{cEQ[3]}{r}", f"={cGE[3]}{r}-{cGE[4]}{r}", fmt=NUM6)
        put(ws, f"{cEQ[4]}{r}", f"={cGE[4]}{r}-{cGE[5]}{r}", fmt=NUM6)
        put(ws, f"{cEQ[5]}{r}", f"={cGE[5]}{r}", fmt=NUM6)
        for k in (3, 4, 5):
            put(ws, f"{cPay[k]}{r}", f"={PAY[s][k]}", F_LINK, "0.0")
        put(ws, f"{cE}{r}", f"=SUMPRODUCT({cEQ[3]}{r}:{cEQ[5]}{r},{cPay[3]}{r}:{cPay[5]}{r})", fmt=NUM6)
        put(ws, f"{cH}{r}", f"=SUMPRODUCT({cEQ[3]}{r}:{cEQ[5]}{r},--({cPay[3]}{r}:{cPay[5]}{r}>0))", fmt=NUM6)
        r += 1
    put(ws, f"A{r}", "TOTAL (one line)", F_B)
    put(ws, f"{cE}{r}", f"=SUM({cE}{top+3}:{cE}{r-1})", fmt=NUM6, bold=True, fill=FILL_KEY)
    put(ws, f"{cH}{r}", f"=SUM({cH}{top+3}:{cH}{r-1})", fmt=NUM6, bold=True, fill=FILL_KEY)
    put(ws, f"A{r+1}", "Validity check: wild count on reel 1 must be 0 for the closed form")
    put(ws, f"{cE}{r+1}", f'=IF({prb["W"][0]}=0,"OK","INVALID: wild on reel 1")', bold=True)
    return f"{ws.title}!${cE}${r}", f"{ws.title}!${cH}${r}", f"{ws.title}!${cE}${r+1}"


wsL = wb.create_sheet("Lines_Base")
put(wsL, "A1", "Kastel Fortune — single-line expected value, base game (exact, closed form)", F_T)
LINE_EV_B, LINE_HIT_B, LINE_OK_B = lines_block(wsL, 3, PRB_B, "Base-game strips", "Reels_Base")
widths(wsL, {"A": 18, **{L(c): 11 for c in range(2, 26)}})
wsF = wb.create_sheet("Lines_FS")
put(wsF, "A1", "Kastel Fortune — single-line expected value, free-spin strips (unmultiplied)", F_T)
LINE_EV_F, LINE_HIT_F, LINE_OK_F = lines_block(wsF, 3, PRB_F, "Free-spin strips", "Reels_FS")
widths(wsF, {"A": 18, **{L(c): 11 for c in range(2, 26)}})

# =============================================================================
# Scatters (Poisson-binomial DP)
# =============================================================================
wsS = wb.create_sheet("Scatters")
put(wsS, "A1", "Kastel Fortune — scatter-count distribution (exact)", F_T)
put(wsS, "A2", "Each strip carries exactly one scatter, so P(scatter visible on reel i) = rows x count / length exactly (the 3 stops "
               "that show it). The total count is a sum of independent Bernoullis — computed by the DP  P_i(k) = P_(i-1)(k)(1-p_i) + P_(i-1)(k-1) p_i.")


def scatter_block(ws, top, cnt, ln, title):
    put(ws, f"A{top}", title, F_B)
    header(ws, top + 1, 1, ["Reel", "S count", "Length", "p_i = rows x count / length"] + [f"P(k={k})" for k in range(0, 6)])
    put(ws, f"A{top+2}", 0, F_B); put(ws, f"E{top+2}", 1.0, F_IN, NUM6)
    for k in range(1, 6):
        put(ws, f"{L(5+k)}{top+2}", 0.0, F_IN, NUM6)
    for i in range(N_REELS):
        r = top + 3 + i
        put(ws, f"A{r}", i + 1, F_B)
        put(ws, f"B{r}", f"={cnt['S'][i]}", F_LINK)
        put(ws, f"C{r}", f"={ln[i]}", F_LINK)
        put(ws, f"D{r}", f"={P['Rows']}*B{r}/C{r}", fmt=NUM6)
        put(ws, f"E{r}", f"=E{r-1}*(1-$D{r})", fmt=NUM6)
        for k in range(1, 6):
            c, cp = L(5 + k), L(4 + k)
            put(ws, f"{c}{r}", f"={c}{r-1}*(1-$D{r})+{cp}{r-1}*$D{r}", fmt=NUM6)
    rr = top + 3 + N_REELS - 1
    put(ws, f"A{rr+1}", "check: sums to 1", F_B); put(ws, f"E{rr+1}", f"=SUM(E{rr}:J{rr})", fmt="0.000000")
    dist = {k: f"{ws.title}!${L(5+k)}${rr}" for k in range(0, 6)}
    return dist, rr


SC_B, rB = scatter_block(wsS, 4, CNT_B, LEN_B, "Base game")
put(wsS, f"A{rB+3}", "Base-game scatter pays and trigger", F_B)
header(wsS, rB + 4, 1, ["Scatters", "P(k)", "1 in", "Pay (x bet)", "RTP share", "Free spins"])
SC_RTP_ROWS = []
for j, k in enumerate((3, 4, 5)):
    r = rB + 5 + j
    put(wsS, f"A{r}", k, F_B); put(wsS, f"B{r}", f"={SC_B[k]}", fmt=SCI); put(wsS, f"C{r}", f"=1/B{r}", fmt="#,##0")
    put(wsS, f"D{r}", f"={SC_PAY[k]}", F_LINK, "0.0"); put(wsS, f"E{r}", f"=B{r}*D{r}", fmt=PCT4); put(wsS, f"F{r}", f"={SC_SPINS[k]}", F_LINK)
    SC_RTP_ROWS.append(r)
rt = rB + 8
put(wsS, f"A{rt}", "P(trigger) = P(3+)", F_B); put(wsS, f"B{rt}", f"=SUM(B{rB+5}:B{rB+7})", fmt=NUM6, fill=FILL_KEY); put(wsS, f"C{rt}", f"=1/B{rt}", fmt="0.0")
put(wsS, f"A{rt+1}", "Scatter RTP", F_B); put(wsS, f"E{rt+1}", f"=SUM(E{rB+5}:E{rB+7})", fmt=PCT4, fill=FILL_KEY)
P_TRIG = f"Scatters!$B${rt}"; SC_RTP = f"Scatters!$E${rt+1}"
SC_F, rF = scatter_block(wsS, rt + 4, CNT_F, LEN_F, "Free spins (retrigger probability per free spin)")
put(wsS, f"A{rF+3}", "P(retrigger per free spin) = P(3+)", F_B); put(wsS, f"B{rF+3}", f"={SC_F[3]}+{SC_F[4]}+{SC_F[5]}", fmt=NUM6, fill=FILL_KEY)
P_RETRIG_F = f"Scatters!$B${rF+3}"
widths(wsS, {"A": 34, "B": 12, "C": 10, "D": 24, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12, "J": 12})

# =============================================================================
# FreeSpins — Markov chain mean recursion
# =============================================================================
wsM = wb.create_sheet("FreeSpins")
put(wsM, "A1", "Kastel Fortune — free spins as an absorbing Markov chain (expected payout, exact)", F_T)
put(wsM, "A2", "State = (spins left n, multiplier m, retriggers used r). E(n,m,r) = m*(mu_wr+mu_wn) + p_wn*E(n-1,m',r) + p_nn*E(n-1,m,r) "
               "+ p_wr*E(n+4,m',r+1) + p_nr*E(n+4,m,r+1); m' = min(m+1,5); when r = 4 the retrigger branches go to (n-1, ., 4). E(0,.,.) = 0.")
put(wsM, "A3", "Per-free-spin outcome classes (joint law of 'line win > 0' and '3+ scatters'); the four probabilities and two means come from the "
               "full enumeration of the free-spin strips (slotmath/exact.py: 12,644,352 combinations). Cross-checks against this workbook's own closed forms are below.")
header(wsM, 5, 1, ["Class", "Probability", "E[line win x 1{class}] (unmultiplied)", "Source"])
cls = [("win & retrigger", "win_retrig"), ("win & no retrigger", "win_noretrig"), ("no win & retrigger", "nowin_retrig"), ("no win & no retrigger", "nowin_noretrig")]
feat = json.loads((ROOT / "results" / "kastel_exact.json").read_text())
cp = feat["fs_class_probs"]
# class means are not in the json: take them from the pickled report
import pickle                                             # noqa: E402
rep = pickle.load(open(ROOT / "results" / "kastel_exact.pkl", "rb"))
cm = rep.feature.class_mean_win
PC, MU = {}, {}
for j, (label, key) in enumerate(cls):
    r = 6 + j
    put(wsM, f"A{r}", label, F_B)
    put(wsM, f"B{r}", float(cp[key]), F_IN, "0.00000000")
    put(wsM, f"C{r}", float(cm[key]), F_IN, "0.00000000")
    put(wsM, f"D{r}", "full enumeration (results/kastel_exact.pkl)")
    PC[key] = f"$B${r}"; MU[key] = f"$C${r}"
put(wsM, "A10", "check: probabilities sum to 1", F_B); put(wsM, "B10", "=SUM(B6:B9)", fmt="0.00000000")
put(wsM, "A11", "check: P(retrigger) vs Scatters sheet (closed form)", F_B); put(wsM, "B11", "=B6+B8", fmt="0.00000000"); put(wsM, "C11", f"={P_RETRIG_F}", F_LINK, "0.00000000"); put(wsM, "D11", '=IF(ABS(B11-C11)<1E-9,"OK","MISMATCH")', bold=True)
put(wsM, "A12", "check: E[line win per free spin] vs Lines_FS (closed form) x lines", F_B); put(wsM, "B12", "=C6+C7", fmt="0.00000000"); put(wsM, "C12", f"={P['Paylines']}*{LINE_EV_F}", F_LINK, "0.00000000"); put(wsM, "D12", '=IF(ABS(B12-C12)<1E-9,"OK","MISMATCH")', bold=True)

G_TOP = 16                                                 # header rows 14 (m), 15 (r); n = 0 at row 16
G_COL0 = 2                                                 # column B = (m=1, r=0)
MULT_CAP, MAX_R, RETRIG = rules.mult_cap, rules.max_retriggers, rules.retrigger_spins
N_MAX = max(rules.spins_by_trigger.values()) + MAX_R * (RETRIG - 1)   # reachable maximum: 20 + 4*4 = 36


def gcol(m, r):
    return L(G_COL0 + 5 * (m - 1) + r)


def gref(n, m, r):
    return f"{gcol(m, r)}{G_TOP + n}"


put(wsM, f"A{G_TOP-3}", "E(n, m, r): expected remaining payout (x bet) from a state. Grid laid out for +5 spins per retrigger, max 4 retriggers, multiplier 1..5 step 1.", F_B)
put(wsM, f"A{G_TOP-2}", "m →", F_B); put(wsM, f"A{G_TOP-1}", "r →", F_B); put(wsM, f"A{G_TOP}", "n = 0", F_B)
for m in range(1, MULT_CAP + 1):
    for r in range(0, MAX_R + 1):
        c = gcol(m, r)
        put(wsM, f"{c}{G_TOP-2}", m, F_H, fill=FILL_H, align="center")
        put(wsM, f"{c}{G_TOP-1}", r, F_H, fill=FILL_H, align="center")
        put(wsM, f"{c}{G_TOP}", 0.0, F_IN, "0.00")
        wsM.column_dimensions[c].width = 9
mu_sum = f"({MU['win_retrig']}+{MU['win_noretrig']})"
for n in range(1, N_MAX + 1):
    put(wsM, f"A{G_TOP+n}", n, F_B)
    for r in range(0, MAX_R + 1):
        if n > max(rules.spins_by_trigger.values()) + (RETRIG - 1) * r:
            continue                                        # unreachable state at this retrigger count
        for m in range(1, MULT_CAP + 1):
            m2 = min(m + rules.mult_step, MULT_CAP)
            if r < MAX_R:
                nR, rR = n - 1 + RETRIG, r + 1
            else:
                nR, rR = n - 1, r
            f = (f"={m}*{mu_sum}"
                 f"+{PC['win_noretrig']}*{gref(n-1, m2, r)}"
                 f"+{PC['nowin_noretrig']}*{gref(n-1, m, r)}"
                 f"+{PC['win_retrig']}*{gref(nR, m2, rR)}"
                 f"+{PC['nowin_retrig']}*{gref(nR, m, rR)}")
            put(wsM, gref(n, m, r), f, fmt="0.00", fill=FILL_GRID)
wsM.column_dimensions["A"].width = 44
# outputs: E[F | k scatters] = E(initial spins, 1, 0) looked up by the spins input on Paytable
ro = G_TOP + N_MAX + 3
put(wsM, f"A{ro}", "Feature value by trigger size", F_B)
header(wsM, ro + 1, 1, ["Scatters", "Initial spins", "E[feature payout] x bet", "P(trigger with k)", "RTP contribution"])
FEAT_EV, FEAT_RTP_ROWS = {}, []
col10 = gcol(1, 0)
for j, k in enumerate((3, 4, 5)):
    r = ro + 2 + j
    put(wsM, f"A{r}", k, F_B); put(wsM, f"B{r}", f"={SC_SPINS[k]}", F_LINK)
    put(wsM, f"C{r}", f"=INDEX(${col10}${G_TOP}:${col10}${G_TOP+N_MAX},B{r}+1)", fmt="0.0000", fill=FILL_KEY)
    put(wsM, f"D{r}", f"={SC_B[k]}", F_LINK, SCI); put(wsM, f"E{r}", f"=C{r}*D{r}", fmt=PCT4)
    FEAT_EV[k] = f"FreeSpins!$C${r}"; FEAT_RTP_ROWS.append(r)
put(wsM, f"A{ro+5}", "Feature RTP (sum)", F_B); put(wsM, f"E{ro+5}", f"=SUM(E{ro+2}:E{ro+4})", fmt=PCT4, fill=FILL_KEY)
FEAT_RTP = f"FreeSpins!$E${ro+5}"
put(wsM, f"A{ro+6}", "Python exact (same recursion, slotmath/markov.py)");
for j, k in enumerate((3, 4, 5)):
    put(wsM, f"{L(2+j)}{ro+6}", float(exact["feature_ev_by_count"][str(k)]), fmt="0.0000", fill=FILL_REF)
put(wsM, f"E{ro+6}", float(exact["rtp_feature"]), fmt=PCT4, fill=FILL_REF)

# =============================================================================
# Dashboard
# =============================================================================
wsD = wb.create_sheet("Dashboard", 0)
put(wsD, "A1", "Kastel Fortune — math model dashboard", F_T)
put(wsD, "A2", "Green = linked from model sheets, black = formula, grey = reference values from the full enumeration in Python (not computable in Excel).")
header(wsD, 4, 1, ["RTP decomposition", "Excel (live)", "Python exact", "Difference"])
rows = [("Line RTP = paylines x E[one line]", f"={P['Paylines']}*{LINE_EV_B}", exact["rtp_line"]),
        ("Scatter-pay RTP", f"={SC_RTP}", exact["rtp_scatter"]),
        ("Free-spins RTP", f"={FEAT_RTP}", exact["rtp_feature"])]
for j, (name, f, ref) in enumerate(rows):
    r = 5 + j
    put(wsD, f"A{r}", name); put(wsD, f"B{r}", f, F_LINK, PCT4); put(wsD, f"C{r}", float(ref), fmt=PCT4, fill=FILL_REF); put(wsD, f"D{r}", f"=B{r}-C{r}", fmt="0.000000%")
put(wsD, "A8", "TOTAL RTP (uncapped)", F_B); put(wsD, "B8", "=SUM(B5:B7)", fmt=PCT4, bold=True, fill=FILL_KEY)
put(wsD, "C8", float(exact["rtp_exact_uncapped"]), fmt=PCT4, fill=FILL_REF); put(wsD, "D8", "=B8-C8", fmt="0.000000%")
put(wsD, "A9", "Max-win cap cost (needs the full distribution)"); put(wsD, "C9", float(exact["cap_cost"]), fmt="0.000000%", fill=FILL_REF)
put(wsD, "A10", "RTP with cap (exact)", F_B); put(wsD, "C10", float(exact["rtp_exact_capped"]), fmt=PCT4, fill=FILL_REF)

header(wsD, 12, 1, ["Frequencies", "Excel (live)", "Python exact", ""])
put(wsD, "A13", "P(one line pays)"); put(wsD, "B13", f"={LINE_HIT_B}", F_LINK, PCT4)
put(wsD, "A14", "Free-spins trigger probability"); put(wsD, "B14", f"={P_TRIG}", F_LINK, PCT4); put(wsD, "C14", float(exact["p_trigger"]), fmt=PCT4, fill=FILL_REF)
put(wsD, "A15", "  ... one feature every N spins"); put(wsD, "B15", "=1/B14", fmt="0.0")
put(wsD, "A16", "E[feature payout | 3 / 4 / 5 scatters]"); put(wsD, "B16", f"={FEAT_EV[3]}", F_LINK, "0.00"); put(wsD, "C16", f"={FEAT_EV[4]}", F_LINK, "0.00"); put(wsD, "D16", f"={FEAT_EV[5]}", F_LINK, "0.00")
put(wsD, "A17", "Closed-form validity (no wild on reel 1)"); put(wsD, "B17", f"={LINE_OK_B}", F_LINK)

header(wsD, 19, 1, ["Needs the full grid enumeration (Python)", "", "Python exact", ""])
refs = [("Hit frequency (any win > 0)", exact["hit_rate"], PCT4), ("Std dev of win per spin (x bet)", exact["std_per_spin"], NUM4),
        ("Volatility index 1.96 x sigma", 1.96 * exact["std_per_spin"], "0.00"),
        ("E[feature | triggered] (x bet)", exact["feature_mean_given_trigger"], "0.00"), ("Median feature payout (x bet)", exact["feature_median_given_trigger"], "0.0"),
        ("P(win >= 100x)  1 in", exact["p_win_ge_100x_1_in"], "#,##0"), ("P(win >= 1000x)  1 in", exact["p_win_ge_1000x_1_in"], "#,##0"),
        ("P(max win 1,500x)  1 in", exact["p_max_win_1_in"], "0.00E+00"),
        ("Base-game stop combinations enumerated", exact["base_combinations"], "#,##0"), ("Free-spin stop combinations enumerated", exact["fs_combinations"], "#,##0")]
for j, (name, v, fmt) in enumerate(refs):
    r = 20 + j
    put(wsD, f"A{r}", name); put(wsD, f"C{r}", float(v), fmt=fmt, fill=FILL_REF)
put(wsD, f"A{20+len(refs)+1}", "Monte-Carlo check: 100,000,000 spins (seed 2024): RTP 96.1923% +/- 0.1480 pp (95% CI), hit 34.815%, sigma 7.5503 — exact value inside the interval (z = -0.13).")
widths(wsD, {"A": 48, "B": 16, "C": 16, "D": 14})

# =============================================================================
# Scenario (what-if on base strips)
# =============================================================================
wsC = wb.create_sheet("Scenario")
put(wsC, "A1", "Kastel Fortune — what-if: edit base-strip symbol COUNTS and read the RTP (feature strips unchanged)", F_T)
put(wsC, "A2", "Blue counts are a copy of the current base strips. Change them (keep one scatter per reel and no wild on reel 1) to see how the dials move.")
header(wsC, 4, 1, ["Symbol"] + [f"Reel {i+1}" for i in range(N_REELS)])
CNT_S, PRB_S, LEN_S = {}, {}, {}
cnt_now = cfg.strips_table()
for j, s in enumerate(SYM):
    r = 5 + j
    put(wsC, f"A{r}", s, F_B); CNT_S[s] = {}
    for i in range(N_REELS):
        put(wsC, f"{L(2+i)}{r}", int(cnt_now[i, SYM.index(s)]), F_IN)
        CNT_S[s][i] = f"Scenario!${L(2+i)}${r}"
rl = 5 + len(SYM)
put(wsC, f"A{rl}", "Strip length", F_B)
for i in range(N_REELS):
    put(wsC, f"{L(2+i)}{rl}", f"=SUM({L(2+i)}5:{L(2+i)}{rl-1})", bold=True); LEN_S[i] = f"Scenario!${L(2+i)}${rl}"
p0 = 2 + N_REELS + 1
header(wsC, 4, p0, ["Symbol"] + [f"p reel {i+1}" for i in range(N_REELS)])
for j, s in enumerate(SYM):
    r = 5 + j
    put(wsC, f"{L(p0)}{r}", s, F_B); PRB_S[s] = {}
    for i in range(N_REELS):
        put(wsC, f"{L(p0+1+i)}{r}", f"={L(2+i)}{r}/{L(2+i)}${rl}", fmt=PCT2); PRB_S[s][i] = f"Scenario!${L(p0+1+i)}${r}"
top = rl + 3
LINE_EV_S, LINE_HIT_S, LINE_OK_S = lines_block(wsC, top, PRB_S, "Scenario line EV (closed form on the scenario counts)", "Scenario")
# scenario scatter block
r_sc = top + 3 + len(PAYING) + 4
dist_S, rS = scatter_block(wsC, r_sc, CNT_S, LEN_S, "Scenario scatter distribution")
ro = rS + 3
header(wsC, ro, 1, ["Result", "Scenario", "Current", "Delta"])
out = [("Line RTP", f"={P['Paylines']}*{LINE_EV_S}", f"=Dashboard!$B$5"),
       ("Scatter-pay RTP", f"={dist_S[3]}*{SC_PAY[3]}+{dist_S[4]}*{SC_PAY[4]}+{dist_S[5]}*{SC_PAY[5]}", f"=Dashboard!$B$6"),
       ("Free-spins RTP (feature strips unchanged)", f"={dist_S[3]}*{FEAT_EV[3]}+{dist_S[4]}*{FEAT_EV[4]}+{dist_S[5]}*{FEAT_EV[5]}", f"=Dashboard!$B$7"),
       ("TOTAL RTP", f"=SUM(B{ro+1}:B{ro+3})", f"=Dashboard!$B$8"),
       ("Feature: 1 in N spins", f"=1/({dist_S[3]}+{dist_S[4]}+{dist_S[5]})", f"=Dashboard!$B$15"),
       ("P(one line pays)", f"={LINE_HIT_S}", f"=Dashboard!$B$13")]
for j, (name, fs, fc) in enumerate(out):
    r = ro + 1 + j
    fmt = "0.0" if "1 in" in name else PCT4
    put(wsC, f"A{r}", name, bold=(name == "TOTAL RTP")); put(wsC, f"B{r}", fs, fmt=fmt, fill=(FILL_KEY if name == "TOTAL RTP" else None))
    put(wsC, f"C{r}", fc, F_LINK, fmt); put(wsC, f"D{r}", f"=B{r}-C{r}", fmt=fmt)
put(wsC, f"A{ro+8}", "Validity: no wild on reel 1"); put(wsC, f"B{ro+8}", f"={LINE_OK_S}", bold=True)
widths(wsC, {"A": 40, **{L(c): 11 for c in range(2, 26)}})

# =============================================================================
# Bets
# =============================================================================
wsB = wb.create_sheet("Bets")
put(wsB, "A1", "Kastel Fortune — bet levels and exposure", F_T)
put(wsB, "A2", "Blue = inputs. Fill the jurisdiction's maximum stake (if any) in B4 to flag bet levels above it.")
put(wsB, "A4", "Jurisdiction max stake (currency, blank = none)", F_B); put(wsB, "B4", None, F_IN, fill=FILL_KEY)
header(wsB, 6, 1, ["Bet level", "Per-line unit (bet / lines)", "Max win (cap x bet)", "Expected return per spin (RTP x bet)", "Expected loss per spin", "Check vs max stake"])
for j, b in enumerate((0.10, 0.20, 0.50, 1, 2, 5, 10, 20, 50, 100)):
    r = 7 + j
    put(wsB, f"A{r}", float(b), F_IN, "0.00"); put(wsB, f"B{r}", f"=A{r}/{P['Paylines']}", fmt="0.0000")
    put(wsB, f"C{r}", f"=A{r}*{P['Max win cap (x bet)']}", fmt="#,##0.00"); put(wsB, f"D{r}", f"=A{r}*Dashboard!$B$8", fmt="0.0000")
    put(wsB, f"E{r}", f"=A{r}-D{r}", fmt="0.0000"); put(wsB, f"F{r}", f'=IF($B$4="","",IF(A{r}>$B$4,"ABOVE MAX STAKE","ok"))')
widths(wsB, {"A": 12, "B": 22, "C": 20, "D": 30, "E": 20, "F": 20})

# =============================================================================
# README
# =============================================================================
wsR = wb.create_sheet("README", 0)
put(wsR, "A1", "Kastel Fortune — mathematical model (Excel edition)", F_T)
lines = [
    "Original 5x3, 20-payline slot with free spins (multiplier ladder x1→x5, retriggers), max win 1,500x. Companion to the Python repository slot-math-portfolio.",
    "",
    "LEGEND:  blue text = hard-coded input you may edit  ·  black = formula  ·  green = link to another sheet  ·  grey fill = reference value computed in Python (full enumeration)  ·  yellow = key output.",
    "",
    "WHAT IS EXACT IN THIS WORKBOOK (pure formulas):",
    "  • Reel-strip counts and per-reel symbol probabilities (COUNTIF / COUNTA) — sheets Reels_Base, Reels_FS.",
    "  • Single-line expected value, closed form P(run>=k) = p1(s)·Π(p_i(s)+p_i(W)) — exact because reel 1 carries no wild — sheets Lines_Base, Lines_FS.",
    "  • Scatter-count distribution (Poisson-binomial dynamic programme) — sheet Scatters.",
    "  • Free-spins expected payout: the absorbing Markov chain's mean recursion over states (spins left, multiplier, retriggers) — sheet FreeSpins.",
    "  • RTP = 20 × E[line] + scatter pays + Σ P(k scatters) × E[feature | k] — sheet Dashboard; the Python exact value is shown next to it with the difference (should be 0).",
    "",
    "WHAT NEEDS THE FULL ENUMERATION (Python, shown as grey reference values): hit frequency, variance / volatility, tail probabilities, max-win reachability, cap cost.",
    "Those depend on the joint law of 20 overlapping lines and are obtained by enumerating all 55,520,640 stop combinations (slotmath/exact.py).",
    "",
    "INPUTS TAKEN FROM PYTHON (sheet FreeSpins, blue): the four per-free-spin outcome-class probabilities and two class means, from the enumeration of the free-spin strips. "
    "Two cross-checks on that sheet recompute what can be recomputed here (retrigger probability, E[line win per free spin]) and must read OK.",
    "",
    "HOW TO USE:  edit a symbol in Reels_Base → every probability, the line EV and the RTP update.  Use sheet Scenario to try symbol-count changes without touching the strips.  Sheet Bets maps bet levels to exposure.",
    "",
    "ASSUMPTIONS:  reels stop uniformly and independently; lines pay left to right, best win per line; scatters pay anywhere and do not pay inside free spins; the multiplier climbs after a winning free spin and applies to the following spins; the max-win cap truncates the total win of a spin including its feature.",
]
for j, t in enumerate(lines):
    put(wsR, f"A{3+j}", t)
wsR.column_dimensions["A"].width = 160

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(f"saved {OUT}")
